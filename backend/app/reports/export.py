"""
Excel / Word / PowerPoint exports in GT house style.

Every metric follows the same sequence as the app: dot plot, box plot iterations,
I-MR chart, basic statistics, histogram, outlier list.

Charts are drawn with matplotlib using the same tokens as the web app so the
documents read as the same product.
"""
from __future__ import annotations

import io

import numpy as np
from datetime import datetime

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from matplotlib.ticker import FuncFormatter, LogLocator, NullFormatter  # noqa: E402
from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Inches, Pt, RGBColor  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor as PRGB  # noqa: E402
from pptx.util import Inches as PI  # noqa: E402
from pptx.util import Pt as PPt  # noqa: E402

BRAND = "4E2A84"
INK = "1B1B24"
MUTED = "6C6C78"
FAINT = "9A9AA4"
LINE = "E7E4EC"
SOFT = "F1ECF8"
SERIES = "#3B5BA5"
RAG = {"green": ("2F7D5B", "E7F3EC"), "amber": ("B7830B", "FBF2DD"), "red": ("BC3A2E", "FBEAE7"), "grey": ("6C6C78", "F0EFF2")}
RAG_TEXT = {"green": "Green", "amber": "Amber", "red": "Red"}

plt.rcParams.update({
    "font.family": "DejaVu Sans", "font.size": 9, "axes.edgecolor": "#" + LINE, "axes.labelcolor": "#" + MUTED,
    "xtick.color": "#" + MUTED, "ytick.color": "#" + MUTED, "axes.spines.top": False, "axes.spines.right": False,
    "axes.grid": True, "grid.color": "#F0EEF3", "grid.linewidth": 0.8, "axes.titleweight": "bold",
    "axes.titlesize": 10, "axes.titlecolor": "#" + INK,
})


def fmt(v, digits: int = 4) -> str:
    if v is None:
        return "—"
    if isinstance(v, (int,)) and not isinstance(v, bool):
        return str(v)
    try:
        f = float(v)
    except (TypeError, ValueError):
        return str(v)
    if f == 0:
        return "0"
    a = abs(f)
    if a >= 100:
        return f"{f:,.1f}"
    if a >= 1:
        return f"{f:.3f}"
    return f"{f:.{digits}f}"


def _ok(metrics: list[dict]) -> list[dict]:
    return [m for m in metrics if "stats" in m]


# ------------------------------------------------------------------ charts --
def control_chart_png(m: dict) -> bytes:
    pts = m["points"]
    seq = m["chart"]["seq"]
    vals = [pts[i]["value"] for i in seq]
    fig, (a1, a2) = plt.subplots(2, 1, figsize=(8, 4.2), gridspec_kw={"height_ratios": [2.2, 1]}, sharex=True)
    xs = list(range(1, len(vals) + 1))
    a1.plot(xs, vals, color=SERIES, lw=1.4, marker="o", ms=3.5, zorder=2)
    for j, i in enumerate(seq):
        if pts[i].get("rules"):
            a1.scatter([j + 1], [vals[j]], s=40, color="#BC3A2E", zorder=3, edgecolor="white", linewidth=1.2)
    c = m["chart"]
    for y, lab, ls in ((c["ucl"], "UCL", "--"), (c["cl"], "CL", "-"), (c["lcl"], "LCL", "--")):
        a1.axhline(y, color="#" + (INK if lab == "CL" else MUTED), lw=1, ls=ls, zorder=1)
        a1.text(1.01, y, f"{lab} {fmt(y)}", va="center", fontsize=7.5, color="#" + MUTED,
                transform=a1.get_yaxis_transform(), clip_on=False)
    if c["scale"] == "log":
        a1.set_yscale("log")
        a1.yaxis.set_major_locator(LogLocator(base=10, subs=(1, 2, 5)))
        a1.yaxis.set_major_formatter(FuncFormatter(lambda v, _: fmt(v)))
        a1.yaxis.set_minor_formatter(NullFormatter())
    a1.set_xlim(0.3, len(vals) + 0.7)
    a1.set_title(f"I-MR chart · {m['name']}", loc="left")
    mr = [v for v in c["mr"]]
    a2.plot(xs, [v if v is not None else float("nan") for v in mr], color=SERIES, lw=1.2, marker="o", ms=3)
    if c["scale"] == "linear":
        a2.axhline(c["mr_bar"], color="#" + INK, lw=1)
        a2.axhline(c["mr_ucl"], color="#" + MUTED, lw=1, ls="--")
    a2.set_xlim(0.3, len(vals) + 0.7)
    a2.set_ylabel("Moving range" + (" (log)" if c["scale"] == "log" else ""), fontsize=8)
    a2.set_xlabel("Observation (time order)", fontsize=8)
    fig.tight_layout()
    return _png(fig)


def histogram_png(m: dict) -> bytes:
    fig, ax = plt.subplots(figsize=(5, 3.2))
    h = m["histogram"]
    for b in h["bins"]:
        ax.bar(b["x0"], b["count"], width=(b["x1"] - b["x0"]) * 0.94, align="edge", color=SERIES, alpha=0.85)
    if h["curve"]:
        ax.plot([p[0] for p in h["curve"]], [p[1] for p in h["curve"]], color="#" + INK, lw=1.4)
    bl = m["baseline"]
    for y, lab in ((bl["lcl"], "LCL"), (bl["cl"], "CL"), (bl["ucl"], "UCL")):
        ax.axvline(y, color="#" + MUTED, lw=1, ls="-" if lab == "CL" else "--")
    ax.set_title(f"Distribution · best fit {m['fit']['best']}", loc="left")
    ax.set_xlabel(m.get("unit") or "")
    fig.tight_layout()
    return _png(fig)


def _png(fig) -> bytes:
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=160)
    plt.close(fig)
    return buf.getvalue()


# ------------------------------------------------------------------- excel --
RED = "#BC3A2E"
AMBER = "#B7830B"
GREY = "#B4B2BC"
STEPS = ["Dot plot", "Box plot", "I-MR chart", "Basic statistics", "Histogram", "Outlier list"]
RULE_TEXT = {1: "beyond 3σ", 2: "9 in a row on one side", 3: "6 in a row trending",
             5: "2 of 3 beyond 2σ", 6: "4 of 5 beyond 1σ"}


# ------------------------------------------------------------- content ----
def outliers_of(m: dict) -> list[dict]:
    """Removed points in removal order: box-plot iterations first, then manual."""
    pts = [p for p in m.get("points", []) if p.get("excluded")]
    pts.sort(key=lambda p: (p.get("removed_by") != "box", p.get("removed_in") or 0, p["row"]))
    return [{"id": p["id"], "value": p["value"],
             "step": f"Box plot, iteration {p['removed_in']}" if p.get("removed_by") == "box" else "Removed by user",
             "reason": p.get("reason") or ""} for p in pts]


def kept_outliers_of(m: dict) -> list[dict]:
    seen, out = set(), []
    for it in m.get("box_iterations", []):
        for k in it["kept"]:
            if k["row"] not in seen:
                seen.add(k["row"])
                out.append(k)
    return out


def signals_of(m: dict) -> list[dict]:
    return [{"id": p["id"], "value": p["value"], "tests": ", ".join(f"Test {r} ({RULE_TEXT.get(r, '')})" for r in p["rules"])}
            for p in m.get("points", []) if p.get("rules") and not p.get("excluded")]


def stats_rows(m: dict) -> list[tuple[str, str]]:
    s, bl = m["stats"], m["baseline"]
    ci = s.get("mean_ci")
    p = s.get("normality_p")
    return [
        ("Data points (n)", f"{m['n_used']} used of {m['n_total']}"),
        ("Mean", fmt(s["mean"])), ("Median", fmt(s["median"])), ("Standard deviation", fmt(s["sd"])),
        ("Minimum", fmt(s["min"])), ("Maximum", fmt(s["max"])),
        ("Q1 (25th percentile)", fmt(s["p25"])), ("Q3 (75th percentile)", fmt(s["p75"])),
        ("95% CI of mean", f"{fmt(ci[0])} – {fmt(ci[1])}" if ci else "—"),
        ("Normality (Anderson-Darling)", "—" if p is None else f"AD {fmt(s.get('normality_ad'), 3)}, p {'< 0.005' if p < 0.005 else f'{p:.3f}'}"
         + (" · normal" if p is not None and p >= 0.05 else " · not normal")),
        ("Best-fit distribution", m["fit"]["best"]),
        ("Baseline LCL", fmt(bl["lcl"])), ("Baseline CL", fmt(bl["cl"])), ("Baseline UCL", fmt(bl["ucl"])),
        ("Limit method", bl["method"]),
        ("Status", f"{RAG_TEXT.get(m['rag'], m['rag'])} · {'stable' if m['status']['stable'] else 'signals present'}"),
    ]


# -------------------------------------------------------------- charts ----
def dotplot_png(m: dict, w: float = 8, h: float = 3.2) -> bytes:
    """Stacked dot plot of all points; removed points marked."""
    pts = m["points"]
    vals = np.array([p["value"] for p in pts])
    nb = min(40, max(12, len(vals)))
    lo, hi = vals.min(), vals.max()
    width = (hi - lo) / nb or 1.0
    stacks: dict[int, int] = {}
    fig, ax = plt.subplots(figsize=(w, h))
    groups = {"inc": ([], []), "box": ([], []), "manual": ([], [])}
    for p in sorted(pts, key=lambda p: p["value"]):
        b = int((p["value"] - lo) / width) if width else 0
        stacks[b] = stacks.get(b, 0) + 1
        key = "inc" if not p.get("excluded") else ("box" if p.get("removed_by") == "box" else "manual")
        groups[key][0].append(lo + (b + 0.5) * width)
        groups[key][1].append(stacks[b])
    ax.scatter(*groups["inc"], s=34, color=SERIES, label="Kept", zorder=3)
    if groups["box"][0]:
        ax.scatter(*groups["box"], s=40, color=RED, label="Removed (box plot)", zorder=3)
    if groups["manual"][0]:
        ax.scatter(*groups["manual"], s=40, facecolor="white", edgecolor=GREY, linewidth=1.5, label="Removed (by user)", zorder=3)
    ax.set_yticks([])
    ax.spines["left"].set_visible(False)
    ax.set_ylim(0, max(stacks.values()) + 1.5)
    ax.set_xlabel(m.get("unit") or "")
    ax.set_title(f"Dot plot · all {len(pts)} points", loc="left")
    ax.legend(loc="upper right", frameon=False, fontsize=8)
    ax.grid(axis="y", visible=False)
    fig.tight_layout()
    return _png(fig)


def boxplots_png(m: dict, w: float = 8) -> bytes:
    """One horizontal box plot per iteration, top to bottom."""
    its = m["box_iterations"]
    allv = [p["value"] for p in m["points"]]
    lo, hi = min(allv), max(allv)
    pad = (hi - lo) * 0.04 or 1
    fig, axes = plt.subplots(len(its), 1, figsize=(w, 0.4 + 1.25 * len(its)), sharex=True, squeeze=False)
    for ax, it in zip(axes[:, 0], its):
        v = np.array(it["values"])
        ax.boxplot(v, vert=False, widths=0.5, whis=1.5, patch_artist=True, showfliers=False,
                   boxprops={"facecolor": "#EEF2FB", "edgecolor": SERIES}, medianprops={"color": SERIES, "linewidth": 2},
                   whiskerprops={"color": "#" + MUTED}, capprops={"color": "#" + MUTED})
        rem = [r["value"] for r in it["removed"]]
        kept = [k["value"] for k in it["kept"]]
        if rem:
            ax.scatter(rem, [1] * len(rem), color=RED, s=36, zorder=3)
        if kept:
            ax.scatter(kept, [1] * len(kept), color=AMBER, s=36, marker="D", zorder=3)
        ax.set_yticks([])
        ax.grid(axis="y", visible=False)
        status = (f"{len(rem)} outlier(s) removed" + (f", {len(kept)} kept" if kept else "") if rem else
                  f"final · {len(kept)} outlier(s) kept" if kept else "no outliers — clean")
        ax.set_title(f"Iteration {it['iteration']} · n = {it['n']} · {status}", loc="left", fontsize=9)
        ax.set_xlim(lo - pad, hi + pad)
    axes[-1, 0].set_xlabel(m.get("unit") or "")
    fig.tight_layout()
    return _png(fig)


# --------------------------------------------------------------- excel ----
def build_xlsx(b: dict) -> bytes:
    wb = Workbook()
    head = Font(bold=True, color="FFFFFF", size=10)
    headfill = PatternFill("solid", fgColor=BRAND)
    thin = Border(bottom=Side(style="thin", color=LINE))
    res = b["results"]

    def table(ws, cols, widths):
        ws.append(cols)
        for c in ws[ws.max_row]:
            c.font, c.fill = head, headfill
        _widths(ws, widths)
        ws.freeze_panes = ws.cell(ws.max_row + 1, 1).coordinate

    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = f"Process Performance Baseline — {b['name']}"
    ws["A1"].font = Font(bold=True, size=14, color=BRAND)
    ws["A2"] = f"Source {b['filename']} · {b['n_rows']} rows · generated {datetime.now():%d %b %Y}"
    ws["A2"].font = Font(color=MUTED)
    ws.append([])
    table(ws, ["Metric", "Unit", "n used", "Removed", "Mean", "Median", "Std dev", "LCL", "CL", "UCL",
               "Limit method", "Best fit", "Stable", "Status"], [38, 12, 8, 9, 10, 10, 10, 10, 10, 10, 34, 13, 8, 9])
    for m in res["metrics"]:
        if "stats" not in m:
            ws.append([m["name"], m.get("unit"), m.get("n_used"), m.get("n_excluded"), m.get("error")])
            continue
        s, bl = m["stats"], m["baseline"]
        ws.append([m["name"], m.get("unit"), m["n_used"], m["n_excluded"], s["mean"], s["median"], s["sd"],
                   bl["lcl"], bl["cl"], bl["ucl"], bl["method"], m["fit"]["best"],
                   "Yes" if m["status"]["stable"] else "No", RAG_TEXT.get(m["rag"], m["rag"])])
        c = ws.cell(ws.max_row, 14)
        fg, bg = RAG.get(m["rag"], RAG["grey"])
        c.fill, c.font = PatternFill("solid", fgColor=bg), Font(color=fg, bold=True)

    ws = wb.create_sheet("Box plot iterations")
    table(ws, ["Metric", "Iteration", "n", "Q1", "Median", "Q3", "Lower fence", "Upper fence", "Removed", "Removed IDs"],
          [38, 9, 6, 10, 10, 10, 12, 12, 9, 50])
    for m in res["metrics"]:
        for it in m.get("box_iterations", []):
            bx = it["box"]
            ws.append([m["name"], it["iteration"], it["n"], bx["q1"], bx["median"], bx["q3"], bx["lower_fence"],
                       bx["upper_fence"], len(it["removed"]), ", ".join(r["id"] for r in it["removed"])])

    ws = wb.create_sheet("Basic statistics")
    labels = None
    for m in _ok(res["metrics"]):
        rows = stats_rows(m)
        if labels is None:
            labels = [k for k, _ in rows]
            table(ws, ["Metric"] + labels, [38] + [16] * len(labels))
        ws.append([m["name"]] + [v for _, v in rows])

    ws = wb.create_sheet("Outlier list")
    table(ws, ["Metric", "ID", "Value", "Step", "Reason"], [38, 16, 12, 22, 70])
    for m in res["metrics"]:
        for o in outliers_of(m):
            ws.append([m["name"], o["id"], o["value"], o["step"], o["reason"]])

    ws = wb.create_sheet("Data points")
    table(ws, ["Metric", "Row", "ID", "Value", "Used", "Removed at", "I-MR signals"], [38, 6, 16, 12, 7, 22, 16])
    for m in res["metrics"]:
        for p in m.get("points", []):
            ws.append([m["name"], p["row"] + 1, p["id"], p["value"], "No" if p["excluded"] else "Yes",
                       (f"Box iteration {p['removed_in']}" if p.get("removed_by") == "box" else
                        "By user" if p.get("removed_by") == "manual" else ""),
                       ", ".join(str(r) for r in p.get("rules") or [])])

    ws = wb.create_sheet("Method")
    for k, v in METHOD:
        ws.append([k, v])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
    _widths(ws, [22, 110])
    for sh in wb.worksheets:
        for row in sh.iter_rows(min_row=2):
            for c in row:
                if isinstance(c.value, float):
                    c.number_format = "0.0000"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


METHOD = [
    ("1. Dot plot", "All data points for the metric, to show spread and extreme values."),
    ("2. Box plot", f"Points beyond 1.5×IQR from the quartiles are removed and the box plot is rebuilt, repeatedly, "
                    f"until no outliers remain or 10% of the points have been removed."),
    ("3. I-MR chart", "Built on the cleaned data, in time order. Points that break the stability tests "
                      "(1, 2, 3, 5, 6) are flagged for investigation, not removed."),
    ("4. Basic statistics", "Descriptive statistics, Anderson-Darling normality test and baseline limits "
                            "(I-MR natural limits for normal data; fitted-distribution percentiles otherwise)."),
    ("5. Histogram", "Distribution of the cleaned data with the best-fitting distribution."),
    ("6. Outlier list", "Every removed point, the step that removed it and the reason."),
]


def _widths(ws, ws_widths: list[int]) -> None:
    for i, w in enumerate(ws_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


# ---------------------------------------------------------------- word ----
def _shade(cell, hex_color: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def _doc_table(doc, header: list[str], rows: list[list[str]], widths: list[float]):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = h
        r = c.paragraphs[0].runs[0]
        r.font.bold, r.font.size, r.font.color.rgb = True, Pt(9), RGBColor(255, 255, 255)
        _shade(c, BRAND)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = str(v)
            cells[i].paragraphs[0].runs[0].font.size = Pt(9)
    for r in t.rows:
        for i, w in enumerate(widths):
            r.cells[i].width = Inches(w)
    return t


def build_docx(b: dict) -> bytes:
    res = b["results"]
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(10)
    for s in ("Heading 1", "Heading 2", "Title"):
        doc.styles[s].font.color.rgb = RGBColor.from_string(BRAND)

    t = doc.add_paragraph()
    r = t.add_run("GRANT THORNTON · PROCESS PERFORMANCE BASELINE")
    r.font.size, r.font.bold, r.font.color.rgb = Pt(9), True, RGBColor.from_string(FAINT)
    doc.add_heading(b["name"], 0)
    meta = doc.add_paragraph()
    meta.add_run(f"Source data: {b['filename']} ({b['n_rows']} rows)\n").bold = True
    meta.add_run(f"Prepared by {b.get('created_by_name') or '—'}. Generated {datetime.now():%d %B %Y}.")

    doc.add_heading("Baseline summary", 1)
    ok = _ok(res["metrics"])
    _doc_table(doc, ["Metric", "n", "LCL", "CL", "UCL", "Fit", "Status"],
               [[m["name"], str(m["n_used"]), fmt(m["baseline"]["lcl"]), fmt(m["baseline"]["cl"]),
                 fmt(m["baseline"]["ucl"]), m["fit"]["best"], RAG_TEXT[m["rag"]]] for m in ok],
               [2.4, 0.5, 0.7, 0.7, 0.7, 1.0, 0.7])
    doc.add_heading("Method", 1)
    for k, v in METHOD:
        para = doc.add_paragraph(style="List Bullet")
        para.add_run(k[3:] + ": ").bold = True
        para.add_run(v)

    for m in ok:
        doc.add_page_break()
        doc.add_heading(m["name"], 1)
        doc.add_paragraph(f"{m.get('unit') or ''} · n = {m['n_used']} of {m['n_total']} · "
                          f"{m['n_box_removed']} removed by box plot, {m['n_manual_removed']} by user · "
                          f"status {RAG_TEXT[m['rag']]}")
        doc.add_heading("1. Dot plot", 2)
        doc.add_picture(io.BytesIO(dotplot_png(m)), width=Inches(6.3))
        doc.add_heading("2. Box plot", 2)
        doc.add_picture(io.BytesIO(boxplots_png(m)), width=Inches(6.3))
        doc.add_heading("3. I-MR chart (after outlier removal)", 2)
        doc.add_picture(io.BytesIO(control_chart_png(m)), width=Inches(6.3))
        sig = signals_of(m)
        doc.add_paragraph("No stability signals." if not sig else
                          "Signals (flagged, not removed): " + "; ".join(f"{s['id']} — {s['tests']}" for s in sig))
        doc.add_heading("4. Basic statistics", 2)
        _doc_table(doc, ["Statistic", "Value"], [list(r) for r in stats_rows(m)], [2.4, 3.9])
        doc.add_heading("5. Histogram", 2)
        doc.add_picture(io.BytesIO(histogram_png(m)), width=Inches(4.6))
        doc.add_heading("6. Outlier list", 2)
        outs = outliers_of(m)
        if outs:
            _doc_table(doc, ["ID", "Value", "Step", "Reason"],
                       [[o["id"], fmt(o["value"]), o["step"], o["reason"]] for o in outs], [1.1, 0.8, 1.5, 2.9])
        else:
            doc.add_paragraph("No points were removed.")
        kept = kept_outliers_of(m)
        if kept:
            doc.add_paragraph("Outliers kept: " + "; ".join(f"{k['id']} ({fmt(k['value'])}, {k['why']})" for k in kept))
        for w in m.get("warnings", []):
            doc.add_paragraph(w, style="List Bullet")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- pptx ----
def _tx(slide, text, x, y, w, h, size=12, bold=False, color=INK, align=None):
    tb = slide.shapes.add_textbox(PI(x), PI(y), PI(w), PI(h))
    tf = tb.text_frame
    tf.word_wrap = True
    para = tf.paragraphs[0]
    para.text = text
    para.font.size, para.font.bold = PPt(size), bold
    para.font.color.rgb = PRGB.from_string(color)
    para.font.name = "Arial"
    if align:
        para.alignment = align
    return tb


def _rect(slide, x, y, w, h, color):
    s = slide.shapes.add_shape(1, PI(x), PI(y), PI(w), PI(h))
    s.fill.solid()
    s.fill.fore_color.rgb = PRGB.from_string(color)
    s.line.fill.background()
    return s


def _bar(slide):
    _rect(slide, 0, 0, 13.333, 0.12, BRAND)


def _content_slide(prs, metric_name: str, step_no: int, title: str):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bar(s)
    _tx(s, f"{metric_name.upper()} · STEP {step_no} OF 6", 0.5, 0.3, 12, 0.3, 10, True, FAINT)
    _tx(s, title, 0.5, 0.55, 12, 0.6, 24, True)
    return s


def _ppt_table(slide, header, rows, x, y, widths, size=10, row_h=0.3):
    tbl = slide.shapes.add_table(len(rows) + 1, len(header), PI(x), PI(y), PI(sum(widths)), PI(row_h * (len(rows) + 1))).table
    for i, w in enumerate(widths):
        tbl.columns[i].width = PI(w)
    for i, h in enumerate(header):
        c = tbl.cell(0, i)
        c.text = h
        c.fill.solid()
        c.fill.fore_color.rgb = PRGB.from_string(BRAND)
        p = c.text_frame.paragraphs[0]
        p.font.size, p.font.bold, p.font.color.rgb = PPt(size), True, PRGB(255, 255, 255)
    for r, row in enumerate(rows, 1):
        for i, v in enumerate(row):
            c = tbl.cell(r, i)
            c.text = str(v)
            c.fill.solid()
            c.fill.fore_color.rgb = PRGB.from_string("FFFFFF" if r % 2 else "F7F6FA")
            p = c.text_frame.paragraphs[0]
            p.font.size = PPt(size - 1)
            p.font.color.rgb = PRGB.from_string(INK)
    return tbl


def build_pptx(b: dict) -> bytes:
    res = b["results"]
    prs = Presentation()
    prs.slide_width, prs.slide_height = PI(13.333), PI(7.5)
    blank = prs.slide_layouts[6]
    ok = _ok(res["metrics"])

    s = prs.slides.add_slide(blank)
    _bar(s)
    _tx(s, "PROCESS PERFORMANCE BASELINE", 0.8, 2.3, 11, 0.4, 12, True, FAINT)
    _tx(s, b["name"], 0.8, 2.7, 11, 1, 36, True, BRAND)
    _tx(s, f"Source {b['filename']} · {b['n_rows']} rows · {len(ok)} metrics", 0.8, 3.8, 11, 0.5, 16, False, MUTED)
    _tx(s, f"Grant Thornton · {datetime.now():%B %Y}", 0.8, 6.6, 11, 0.4, 11, False, FAINT)

    s = prs.slides.add_slide(blank)
    _bar(s)
    _tx(s, "Approach for every metric", 0.5, 0.35, 12, 0.6, 24, True)
    for i, (k, v) in enumerate(METHOD):
        y = 1.35 + i * 0.95
        _rect(s, 0.6, y, 0.55, 0.55, BRAND)
        _tx(s, str(i + 1), 0.6, y + 0.08, 0.55, 0.4, 16, True, "FFFFFF", align=2)
        _tx(s, k[3:], 1.4, y - 0.02, 3, 0.4, 14, True)
        _tx(s, v, 1.4, y + 0.3, 11.3, 0.6, 11, False, MUTED)

    for start in range(0, len(ok), 14):
        chunk = ok[start:start + 14]
        s = prs.slides.add_slide(blank)
        _bar(s)
        _tx(s, "Baseline summary" + (" (continued)" if start else ""), 0.5, 0.35, 12, 0.6, 24, True)
        tbl = _ppt_table(s, ["Metric", "n", "LCL", "CL", "UCL", "Fit", "Status"],
                         [[m["name"], m["n_used"], fmt(m["baseline"]["lcl"]), fmt(m["baseline"]["cl"]),
                           fmt(m["baseline"]["ucl"]), m["fit"]["best"], RAG_TEXT[m["rag"]]] for m in chunk],
                         0.5, 1.2, [5.0, 0.8, 1.3, 1.3, 1.3, 1.6, 1.0])
        for r, m in enumerate(chunk, 1):
            c = tbl.cell(r, 6)
            c.fill.fore_color.rgb = PRGB.from_string(RAG[m["rag"]][1])
            c.text_frame.paragraphs[0].font.color.rgb = PRGB.from_string(RAG[m["rag"]][0])

    for idx, m in enumerate(ok, 1):
        bl, ss = m["baseline"], m["stats"]
        # section divider
        s = prs.slides.add_slide(blank)
        _rect(s, 0, 0, 13.333, 7.5, BRAND)
        _tx(s, f"METRIC {idx} OF {len(ok)}", 0.9, 2.4, 11, 0.4, 13, True, "D9CCEF")
        _tx(s, m["name"], 0.9, 2.8, 11.5, 1.2, 38, True, "FFFFFF")
        _tx(s, f"{m.get('unit') or ''} · n = {m['n_used']} of {m['n_total']} · baseline {fmt(bl['lcl'])} / "
               f"{fmt(bl['cl'])} / {fmt(bl['ucl'])} · {RAG_TEXT[m['rag']]}", 0.9, 4.1, 11.5, 0.5, 15, False, "E9E1F5")
        _tx(s, "  →  ".join(STEPS), 0.9, 6.5, 11.5, 0.4, 12, False, "D9CCEF")

        s = _content_slide(prs, m["name"], 1, "Dot plot")
        s.shapes.add_picture(io.BytesIO(dotplot_png(m, 11, 4.4)), PI(0.5), PI(1.35), width=PI(12.3))
        _tx(s, f"All {m['n_total']} points. Red points are removed by the box plot in the next step.",
            0.5, 6.8, 12.3, 0.4, 11, False, MUTED)

        s = _content_slide(prs, m["name"], 2, f"Box plot · {len(m['box_iterations'])} iteration(s)")
        fig_h = 0.4 + 1.25 * len(m["box_iterations"])
        scale = min(8.9 / 8.6, 5.9 / fig_h)
        s.shapes.add_picture(io.BytesIO(boxplots_png(m, 8.6)), PI(0.4), PI(1.35), width=PI(8.6 * scale))
        _ppt_table(s, ["Iteration", "n", "Removed"],
                   [[it["iteration"], it["n"], len(it["removed"])] for it in m["box_iterations"]],
                   9.6, 1.45, [1.2, 0.8, 1.2], 11)
        _tx(s, f"Removed {m['n_box_removed']} point(s); limit is 10% of the data. Rebuilt until no point lies "
               "beyond 1.5×IQR.", 9.6, 1.6 + 0.3 * (len(m["box_iterations"]) + 1), 3.3, 1.4, 10, False, MUTED)

        s = _content_slide(prs, m["name"], 3, "I-MR chart (after outlier removal)")
        s.shapes.add_picture(io.BytesIO(control_chart_png(m)), PI(0.4), PI(1.35), width=PI(8.6))
        y = 1.5
        for k, v in [("UCL", bl["ucl"]), ("CL", bl["cl"]), ("LCL", bl["lcl"])]:
            _tx(s, k, 9.4, y, 1.2, 0.3, 10, True, FAINT)
            _tx(s, fmt(v), 10.5, y - 0.07, 2.4, 0.4, 18, True)
            y += 0.6
        sig = signals_of(m)
        _tx(s, "Stable: no stability signals." if not sig else
            f"{len(sig)} signal(s), flagged not removed: " + "; ".join(f"{x['id']} ({x['tests']})" for x in sig[:4])
            + ("…" if len(sig) > 4 else ""), 9.4, y + 0.2, 3.6, 2.5, 10, False, MUTED)

        s = _content_slide(prs, m["name"], 4, "Basic statistics")
        rows = stats_rows(m)
        half = (len(rows) + 1) // 2
        _ppt_table(s, ["Statistic", "Value"], [list(r) for r in rows[:half]], 0.5, 1.4, [2.8, 3.2], 11, 0.36)
        _ppt_table(s, ["Statistic", "Value"], [list(r) for r in rows[half:]], 6.8, 1.4, [2.6, 3.4], 11, 0.36)

        s = _content_slide(prs, m["name"], 5, f"Histogram · best fit {m['fit']['best']}")
        s.shapes.add_picture(io.BytesIO(histogram_png(m)), PI(0.5), PI(1.35), height=PI(5.3))
        _ppt_table(s, ["Distribution", "AD", "p-value"],
                   [[f["name"] + (" ✓" if f["name"] == m["fit"]["best"] else ""),
                     "—" if f["ad"] is None else f"{f['ad']:.3f}", f["p_label"]] for f in m["fit"]["fits"]],
                   9.0, 1.5, [1.7, 1.0, 1.2], 11)

        outs = outliers_of(m)
        chunks = [outs[i:i + 12] for i in range(0, len(outs), 12)] or [[]]
        for ci, ch in enumerate(chunks):
            s = _content_slide(prs, m["name"], 6, "Outlier list" + (" (continued)" if ci else ""))
            if ch:
                _ppt_table(s, ["ID", "Value", "Step", "Reason"],
                           [[o["id"], fmt(o["value"]), o["step"], o["reason"][:90]] for o in ch],
                           0.5, 1.4, [2.0, 1.4, 2.4, 6.5], 11, 0.36)
            else:
                _tx(s, "No points were removed for this metric.", 0.5, 1.5, 12, 0.5, 14, False, MUTED)
            kept = kept_outliers_of(m)
            if kept and ci == len(chunks) - 1:
                _tx(s, "Outliers kept: " + "; ".join(f"{k['id']} ({k['why']})" for k in kept),
                    0.5, 6.6, 12.3, 0.6, 10, False, MUTED)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def build(fmt_: str, run: dict) -> bytes:
    return {"xlsx": build_xlsx, "docx": build_docx, "pptx": build_pptx}[fmt_](run)
