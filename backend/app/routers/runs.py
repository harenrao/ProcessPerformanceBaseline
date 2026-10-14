"""PPB runs: upload a file, get the baseline, revisit past runs."""
from __future__ import annotations

import io
import json
import re
from typing import Optional

import pandas as pd
from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from pydantic import BaseModel

from .. import db
from ..config import MAX_UPLOAD_MB
from ..engine.baseline import DEFAULT_SETTINGS, run_baseline
from ..engine.ingest import infer_mapping, propose_metrics, quality_report, read_table
from ..reports import export
from ..security import current_user

router = APIRouter(prefix="/api", tags=["runs"])

# Input layout the tool expects (one row per completed requirement / project).
TEMPLATE_COLUMNS = [
    "Sr.No.", "Requirement ID", "Initial Project Size",
    "Actual Analysis Efforts (MD)", "Actual Analysis Review Efforts (MD)",
    "Design Efforts (MD)", "Design Review Efforts (MD)", "Coding Efforts", "Code Review Efforts",
    "Unit Test Design Efforts", "Unit Test Review Efforts", "SIT Testing Efforts", "UAT Efforts",
    "Documentation Efforts", "Documentation Review Efforts",
    "Analysis Review Defects", "Design Review Defects", "Code Review Defects", "SIT Defects", "UAT Defects",
]

MEDIA = {
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
}


@router.get("/template")
def template(_: dict = Depends(current_user)):
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(TEMPLATE_COLUMNS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="4E2A84")
        ws.column_dimensions[c.column_letter].width = max(14, len(c.value) + 2)
    notes = wb.create_sheet("How to fill")
    for line in [
        "One row per completed requirement, change request or project.",
        "Initial Project Size: in your size unit (e.g. function points). Rows with no size are skipped.",
        "Effort columns: person-days (MD). Leave 0 or blank if a phase did not apply.",
        "Defect columns: number of defects found in that phase. 0 is valid data.",
        "Keep rows in completion order (oldest first) so the control charts are meaningful.",
        "Extra columns are allowed; you can add or rename phases. Headers containing 'effort' or 'defect' are picked up.",
        "Aim for at least 20 rows.",
    ]:
        notes.append([line])
    notes.column_dimensions["A"].width = 110
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), media_type=MEDIA["xlsx"],
                    headers={"Content-Disposition": 'attachment; filename="PPB_input_template.xlsx"'})


async def _read(file: UploadFile) -> tuple[pd.DataFrame, str]:
    content = await file.read()
    if len(content) > MAX_UPLOAD_MB * 1024 * 1024:
        raise HTTPException(413, f"File is larger than {MAX_UPLOAD_MB} MB.")
    name = file.filename or "upload.xlsx"
    try:
        df = read_table(content, name)
    except Exception as e:  # noqa: BLE001
        raise HTTPException(400, f"Could not read the file: {e}")
    if df.empty:
        raise HTTPException(400, "The file has no data rows.")
    return df, name


@router.post("/inspect")
async def inspect(file: UploadFile = File(...), _: dict = Depends(current_user)):
    """Check an upload before running: detected columns and data-quality notes."""
    df, name = await _read(file)
    mapping = infer_mapping(df)
    return {"filename": name, "rows": len(df), "columns": list(df.columns), "mapping": mapping,
            "quality": quality_report(df, mapping, propose_metrics(mapping))}


def _rows(df: pd.DataFrame) -> list[dict]:
    rows = df.astype(object).where(df.notna(), None).to_dict("records")
    for r in rows:
        for k, v in r.items():
            if hasattr(v, "isoformat"):
                r[k] = v.isoformat()
    return rows


def _compute(run: dict) -> dict:
    df = pd.DataFrame(run["rows"], columns=run["columns"])
    return run_baseline(df, run["mapping"], run["metrics"], run["settings"], run["exclusions"])


@router.post("/runs")
async def create_run(file: UploadFile = File(...), name: str = Form(""), size_unit: str = Form("FP"),
                     mapping: Optional[str] = Form(None), u: dict = Depends(current_user)):
    df, fname = await _read(file)
    m = infer_mapping(df)
    if mapping:
        try:
            m.update({k: v for k, v in json.loads(mapping).items() if k in m})
        except ValueError:
            raise HTTPException(400, "Invalid column mapping")
    cols = set(df.columns)
    if not m.get("size") or m["size"] not in cols:
        raise HTTPException(400, "Choose the column that holds project size.")
    m["effort"] = [c for c in m.get("effort", []) if c in cols]
    m["defect"] = [c for c in m.get("defect", []) if c in cols]
    if not m["effort"] and not m["defect"]:
        raise HTTPException(400, "No effort or defect columns found. Check the file against the template.")
    metrics = propose_metrics(m, size_unit.strip() or "FP")
    run = {"name": name.strip() or re.sub(r"\.\w+$", "", fname), "filename": fname, "columns": list(df.columns),
           "rows": _rows(df), "mapping": m, "metrics": metrics, "settings": dict(DEFAULT_SETTINGS), "exclusions": {}}
    run["results"] = _compute(run)
    rid = db.exec_("""INSERT INTO runs (user_id, name, filename, created_at, updated_at, columns, rows, mapping,
                      metrics, settings, exclusions, results) VALUES (?,?,?,?,?,?,?,?,?,?,?,?)""",
                   u["id"], run["name"], fname, db.now(), db.now(), run["columns"], run["rows"], m, metrics,
                   run["settings"], {}, run["results"])
    return get_run(rid, u)


def _load(rid: int, u: dict) -> dict:
    r = db.one("SELECT r.*, u.name AS created_by_name FROM runs r JOIN users u ON u.id = r.user_id "
               "WHERE r.id = ? AND r.user_id = ?", rid, u["id"])
    if not r:
        raise HTTPException(404, "Run not found")
    r["n_rows"] = len(r["rows"])
    return r


@router.get("/runs")
def list_runs(u: dict = Depends(current_user)):
    out = []
    for r in db.all_("SELECT id, name, filename, created_at, updated_at, results FROM runs "
                     "WHERE user_id = ? ORDER BY id DESC", u["id"]):
        s = (r.pop("results") or {}).get("summary") or {}
        out.append({**r, "summary": s})
    return out


@router.get("/runs/{rid}")
def get_run(rid: int, u: dict = Depends(current_user)):
    r = _load(rid, u)
    r.pop("rows")
    return r


class OverrideIn(BaseModel):
    metric: str
    row: int
    action: str  # "exclude" (remove by hand), "keep" (put back a box-plot outlier) or "reset"
    reason: Optional[str] = None


@router.patch("/runs/{rid}/overrides")
def override(rid: int, body: OverrideIn, u: dict = Depends(current_user)):
    r = _load(rid, u)
    ov = r["exclusions"].setdefault(body.metric, {})
    if body.action == "reset":
        ov.pop(str(body.row), None)
    elif body.action in ("exclude", "keep"):
        reason = (body.reason or "").strip()
        if len(reason) < 10:
            raise HTTPException(400, "Describe the reason (at least 10 characters).")
        ov[str(body.row)] = {"action": body.action, "reason": f"{reason} — {u['name']}, {db.now()[:10]}"}
    else:
        raise HTTPException(400, "Unknown action")
    r["results"] = _compute(r)
    db.exec_("UPDATE runs SET exclusions = ?, results = ?, updated_at = ? WHERE id = ?",
             r["exclusions"], r["results"], db.now(), rid)
    return get_run(rid, u)


class RenameIn(BaseModel):
    name: str


@router.patch("/runs/{rid}")
def rename(rid: int, body: RenameIn, u: dict = Depends(current_user)):
    _load(rid, u)
    if not body.name.strip():
        raise HTTPException(400, "Name cannot be empty")
    db.exec_("UPDATE runs SET name = ? WHERE id = ?", body.name.strip(), rid)
    return get_run(rid, u)


@router.delete("/runs/{rid}")
def delete(rid: int, u: dict = Depends(current_user)):
    _load(rid, u)
    db.exec_("DELETE FROM runs WHERE id = ?", rid)
    return {"ok": True}


@router.get("/runs/{rid}/export/{fmt}")
def export_file(rid: int, fmt: str, metrics: Optional[str] = None, u: dict = Depends(current_user)):
    """`metrics` is a comma-separated list of metric keys; omitted means all metrics."""
    if fmt not in MEDIA:
        raise HTTPException(400, "Format must be xlsx, docx or pptx")
    r = _load(rid, u)
    if metrics:
        keys = set(metrics.split(","))
        r["results"]["metrics"] = [m for m in r["results"]["metrics"] if m["key"] in keys]
        if not r["results"]["metrics"]:
            raise HTTPException(400, "Select at least one metric")
    safe = re.sub(r"[^A-Za-z0-9]+", "_", r["name"]).strip("_") or "PPB"
    return Response(export.build(fmt, r), media_type=MEDIA[fmt],
                    headers={"Content-Disposition": f'attachment; filename="PPB_{safe}.{fmt}"'})
